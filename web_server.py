import os
import sys
import datetime
import json
import sqlite3
import tempfile
import math
import logging
import io
import base64
import qrcode
import qrcode.image.svg
from fastapi import FastAPI, Request, Form, UploadFile, File, Query
from fastapi.responses import HTMLResponse, RedirectResponse, FileResponse, JSONResponse
from fastapi.templating import Jinja2Templates

app = FastAPI(title="IT 资产管理系统行政自动化资产大盘", version="2.0")

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

# ── 日志配置 ──────────────────────────────────────────────────────────────────
LOG_DIR = "/tmp/asset_server_logs"
os.makedirs(LOG_DIR, exist_ok=True)
LOG_FILE = os.path.join(LOG_DIR, f"server_{datetime.datetime.now().strftime('%Y%m%d')}.log")

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    handlers=[
        logging.FileHandler(LOG_FILE, encoding="utf-8"),
        logging.StreamHandler(sys.stdout),
    ]
)
log = logging.getLogger("asset_server")
templates = Jinja2Templates(directory=os.path.join(BASE_DIR, "templates"))
DB_PATH = os.path.join(os.path.dirname(__file__), 'data', 'assets.db')
PER_PAGE = 20

def get_db():
    """获取数据库连接，带重试逻辑"""
    max_retries = 3
    for attempt in range(max_retries):
        try:
            conn = sqlite3.connect(DB_PATH, timeout=10)
            conn.row_factory = sqlite3.Row
            try:
                conn.execute("PRAGMA journal_mode=WAL")  # 启用WAL模式，提升并发性能
            except sqlite3.OperationalError:
                pass  # 某些环境可能不支持WAL
            return conn
        except sqlite3.OperationalError as e:
            log.warning(f"DB连接失败 (尝试 {attempt+1}/{max_retries}): {e}")
            if attempt == max_retries - 1:
                raise
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


# ── 健康检查 ───────────────────────────────────────────────────────────────────
@app.get("/health")
async def health_check():
    """健康检查接口：验证服务器和数据库是否正常"""
    try:
        conn = get_db()
        conn.execute("SELECT 1").fetchone()
        conn.close()
        return {"status": "ok", "db": "connected", "timestamp": datetime.datetime.now().isoformat()}
    except Exception as e:
        log.error(f"健康检查失败: {e}")
        return JSONResponse(status_code=503, content={"status": "error", "detail": str(e)})


# ── 调试接口：查看数据库列结构 ─────────────────────────────────────────────────
@app.get("/debug/columns")
async def debug_columns():
    """查看数据库 assets 表的所有列名和一条示例数据"""
    try:
        conn = get_db()
        cols = conn.execute("PRAGMA table_info(assets)").fetchall()
        col_list = [{"cid": c["cid"], "name": c["name"], "type": c["type"]} for c in cols]
        row = conn.execute("SELECT * FROM assets WHERE asset_id = 'DEMO-ASSET-0001'").fetchone()  # 占位示例编码，非真实资产
        sample = dict(row) if row else {"error": "asset not found"}
        conn.close()
        return {"columns": col_list, "column_count": len(col_list), "sample_data": sample}
    except Exception as e:
        log.error(f"调试接口失败: {e}")
        return JSONResponse(status_code=500, content={"error": str(e)})


# ── 全局异常捕获 ───────────────────────────────────────────────────────────────
@app.exception_handler(Exception)
async def global_exception_handler(request: Request, exc: Exception):
    log.error(f"未捕获异常 [{request.method} {request.url}]: {exc}", exc_info=True)
    return HTMLResponse(
        content=f"""<html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"></head>
        <body style="font-family:sans-serif;text-align:center;padding:40px;">
            <h2 style="color:#e74c3c;">⚠️ 服务器内部错误</h2>
            <p style="color:#666;">错误信息已记录，请联系管理员。</p>
            <p style="font-size:12px;color:#999;margin-top:20px;">{datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")}</p>
            <a href="/" style="display:inline-block;margin-top:20px;padding:8px 20px;background:#3498db;color:white;text-decoration:none;border-radius:4px;">返回首页</a>
        </body></html>""",
        status_code=500
    )

status_zh = {"in_use": "分配使用/在用", "in_stock": "闲置入库", "scrapped": "已报废", "disposed": "已处置"}
type_map_zh = {
    'laptop': '笔记本', 'desktop': '台式机', 'monitor': '显示器',
    'phone': '手机', 'tablet': '平板', 'keyboard_mouse': '键鼠',
    'printer': '打印机', 'other': '其他'
}

@app.get("/", response_class=HTMLResponse)
async def dashboard(
    request: Request,
    page: int = Query(1, ge=1),
    status_filter: str = Query("", alias="status"),
    search: str = Query(""),
    region: str = Query("")
):
    log.info(f"访问首页: page={page}, status={status_filter}, search={search}, region={region}")
    conn = None
    try:
        conn = get_db()

        # 构建查询条件
        where_clauses = []
        params = []

        if status_filter:
            where_clauses.append("a.status = ?")
            params.append(status_filter)

        if search:
            where_clauses.append("(a.asset_id LIKE ? OR a.serial_number LIKE ? OR a.assignee LIKE ? OR a.department LIKE ? OR a.model LIKE ? OR a.brand LIKE ? OR a.name LIKE ? OR a.cpu LIKE ? OR a.memory LIKE ? OR a.disk LIKE ?)")
            like = f"%{search}%"
            params.extend([like] * 10)

        if region:
            where_clauses.append("a.location LIKE ?")
            params.append(f"%{region}%")

        where_sql = " WHERE " + " AND ".join(where_clauses) if where_clauses else ""

        # 总数量
        c = conn.execute(f"SELECT count(*) FROM assets a{where_sql}", params)
        total_assets = c.fetchone()[0]
        total_pages = max(1, math.ceil(total_assets / PER_PAGE))
        page = min(page, total_pages)
        offset = (page - 1) * PER_PAGE

        # 分页数据
        c = conn.execute(f"SELECT a.* FROM assets a{where_sql} ORDER BY REPLACE(a.purchase_date,'/','-') DESC, a.asset_id DESC LIMIT ? OFFSET ?", params + [PER_PAGE, offset])
        assets = c.fetchall()

        # 统计数据
        c = conn.execute(f"SELECT count(*) as total, sum(a.purchase_price) as total_price FROM assets a{where_sql}", params)
        stats = c.fetchone()

        # 状态分布（全量）
        c = conn.execute("SELECT status, count(*) as count FROM assets GROUP BY status")
        status_counts = c.fetchall()

        # 价值分布
        c = conn.execute("SELECT asset_type, sum(purchase_price) as total_value FROM assets GROUP BY asset_type")
        type_values = c.fetchall()

        # 地区列表
        c = conn.execute("SELECT DISTINCT location FROM assets WHERE location IS NOT NULL AND location != '' ORDER BY location")
        regions = [r["location"] for r in c.fetchall() if len(r["location"].strip()) <= 20]

        status_data = [{"name": status_zh.get(r["status"], r["status"]), "value": r["count"]} for r in status_counts]
        type_value_data = [{"name": type_map_zh.get(r["asset_type"], r["asset_type"] or "未知"), "value": round(r["total_value"] or 0, 2)} for r in type_values]

        counts = {"in_use": 0, "in_stock": 0, "scrapped": 0, "disposed": 0, "total": stats["total"] if stats else 0}
        for r in status_counts:
            counts[r["status"]] = r["count"]

        # 分页范围
        page_range = []
        if total_pages <= 7:
            page_range = list(range(1, total_pages + 1))
        else:
            page_range = [1]
            start = max(2, page - 2)
            end = min(total_pages - 1, page + 2)
            if start > 2:
                page_range.append("...")
            page_range.extend(range(start, end + 1))
            if end < total_pages - 1:
                page_range.append("...")
            page_range.append(total_pages)

        return templates.TemplateResponse(request, "index.html", {
            "request": request,
            "assets": [dict(a) for a in assets],
            "total_count": total_assets,
            "total_price": round(stats["total_price"], 2) if stats and stats["total_price"] else 0,
            "status_data": json.dumps(status_data, ensure_ascii=False),
            "type_value_data": json.dumps(type_value_data, ensure_ascii=False),
            "counts": counts,
            "type_map_zh": type_map_zh,
            "status_zh": status_zh,
            "page": page,
            "total_pages": total_pages,
            "page_range": page_range,
            "per_page": PER_PAGE,
            "current_status": status_filter,
            "current_search": search,
            "regions": regions,
            "current_region": region
        })
    except Exception as e:
        log.error(f"首页渲染失败: {e}", exc_info=True)
        return HTMLResponse(
            content=f"""<html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"></head>
            <body style="font-family:sans-serif;text-align:center;padding:40px;">
                <h2 style="color:#e74c3c;">⚠️ 数据加载失败</h2>
                <p style="color:#666;">请稍后刷新页面，或联系管理员查看日志。</p>
                <p style="font-size:12px;color:#999;">{datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")}</p>
                <a href="/" style="display:inline-block;margin-top:20px;padding:8px 20px;background:#3498db;color:white;text-decoration:none;border-radius:4px;">重新加载</a>
            </body></html>""",
            status_code=500
        )
    finally:
        if conn:
            try:
                conn.close()
            except:
                pass

@app.post("/api/update_status")
async def update_status(request: Request, asset_id: str = Form(...), new_status: str = Form(...), operator: str = Form("猫猫")):
    conn = None
    try:
        conn = get_db()
        cursor = conn.execute("SELECT status FROM assets WHERE asset_id = ?", (asset_id,))
        row = cursor.fetchone()
        if not row:
            return RedirectResponse(url="/?error=not_found", status_code=303)
            
        old_status = row["status"]
        
        if old_status != new_status:
            timestamp = datetime.datetime.now().isoformat()
            
            # 分配使用 → 自动记录领用时间
            if new_status == "in_use" and old_status != "in_use":
                conn.execute("UPDATE assets SET status = ?, updated_at = ?, assign_date = ? WHERE asset_id = ?", (new_status, timestamp, timestamp, asset_id))
            # 其他状态变更
            else:
                conn.execute("UPDATE assets SET status = ?, updated_at = ? WHERE asset_id = ?", (new_status, timestamp, asset_id))
            
            # 业务规则：闲置入库 → 自动清空使用人 + 记录"退库入库"日志
            if new_status == "in_stock":
                conn.execute("UPDATE assets SET assignee = '' WHERE asset_id = ? AND assignee IS NOT NULL AND assignee != ''", (asset_id,))
                log.info(f"闲置入库自动清空使用人: {asset_id}")
                details = json.dumps({
                    "action": "退库入库",
                    "from_status": old_status,
                    "to_status": "in_stock",
                    "assignee_cleared": True
                }, ensure_ascii=False)
                conn.execute('''
                    INSERT INTO operation_log (timestamp, asset_id, operation, operator, details, old_value, new_value)
                    VALUES (?, ?, ?, ?, ?, ?, ?)
                ''', (timestamp, asset_id, "退库入库", operator, details, old_status, "in_stock"))
            else:
                details = json.dumps({"action": "Web状态修改", "field": "status", "from": old_status, "to": new_status}, ensure_ascii=False)
                conn.execute('''
                    INSERT INTO operation_log (timestamp, asset_id, operation, operator, details, old_value, new_value)
                    VALUES (?, ?, ?, ?, ?, ?, ?)
                ''', (timestamp, asset_id, "STATUS_CHANGE", operator, details, old_status, new_status))
            conn.commit()
            log.info(f"状态修改: {asset_id} {old_status}→{new_status} by {operator}")
    except Exception as e:
        log.error(f"状态修改失败 [{asset_id}]: {e}")
    finally:
        if conn:
            try:
                conn.close()
            except:
                pass
    return RedirectResponse(url="/", status_code=303)


@app.post("/api/assign")
async def assign_asset(request: Request, asset_id: str = Form(...), assignee: str = Form(...), operator: str = Form("猫猫(Web大盘)"), return_to: str = Form("")):
    """派发 / 改派：将资产绑定使用人并流转为在用(in_use)，写入审计日志。
    覆盖场景：闲置入库(in_stock)后重新派发、在用人变更(改派)、待退库后派发。"""
    conn = None
    try:
        conn = get_db()
        cursor = conn.execute("SELECT assignee, status FROM assets WHERE asset_id = ?", (asset_id,))
        row = cursor.fetchone()
        if not row:
            return RedirectResponse(url="/?error=not_found", status_code=303)

        old_assignee = row["assignee"]
        old_status = row["status"]
        timestamp = datetime.datetime.now().isoformat()
        new_assignee = assignee.strip()

        if not new_assignee:
            return RedirectResponse(url="/?error=empty_assignee", status_code=303)

        # 派发：绑定使用人 + 状态转在用 + 记录领用时间（便于审计：领用/派发时间留痕）
        conn.execute(
            "UPDATE assets SET assignee=?, status='in_use', updated_at=?, assign_date=? WHERE asset_id=?",
            (new_assignee, timestamp, timestamp, asset_id)
        )

        details = json.dumps({
            "action": "Web派发分配",
            "assignee": new_assignee,
            "from_status": old_status,
            "from_assignee": old_assignee
        }, ensure_ascii=False)
        conn.execute('''
            INSERT INTO operation_log (timestamp, asset_id, operation, operator, details, old_value, new_value)
            VALUES (?, ?, ?, ?, ?, ?, ?)
        ''', (timestamp, asset_id, "ASSIGN", operator, details,
              f"assignee={old_assignee!r}, status={old_status!r}",
              f"assignee={new_assignee!r}, status='in_use'"))
        conn.commit()
        log.info(f"派发分配: {asset_id} -> {new_assignee} by {operator}")
    except Exception as e:
        log.error(f"派发失败 [{asset_id}]: {e}")
    finally:
        if conn:
            try:
                conn.close()
            except:
                pass
    return RedirectResponse(url=return_to or "/", status_code=303)


def parse_excel_and_insert(tmp_path, filename="导入文件"):
    """解析Excel并插入数据，懒导入openpyxl"""
    try:
        import openpyxl
    except ImportError:
        log.error("openpyxl 未安装，无法解析Excel")
        raise RuntimeError("请先安装 openpyxl：pip install openpyxl")
    
    conn = get_db()
    wb = openpyxl.load_workbook(tmp_path, read_only=True, data_only=True)
    imported_count = 0
    timestamp = datetime.datetime.now().isoformat()
    
    for sheet_name in wb.sheetnames:
        if '在用' in sheet_name:
            status = 'in_use'
        elif '空闲' in sheet_name:
            status = 'in_stock'
        else:
            continue
            
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        for row in rows:
            if not row or len(row) < 2 or not row[1]: 
                continue
                
            asset_id = str(row[1]).strip()
            name = str(row[2]).strip() if len(row) > 2 and row[2] else ''
            
            asset_type_raw = str(row[3]).strip() if len(row) > 3 and row[3] else ''
            asset_type = 'other'
            TYPE_MAP = {
                '台式机': 'desktop', '笔记本': 'laptop', '显示器': 'monitor',
                '手机': 'phone', '平板': 'tablet', '键鼠': 'keyboard_mouse',
                '打印机': 'printer', '其他': 'other',
                '一体机': 'desktop', '电脑主机': 'desktop'
            }
            for k, v in TYPE_MAP.items():
                if k in asset_type_raw:
                    asset_type = v
                    break
                    
            brand = str(row[4]).strip() if len(row) > 4 and row[4] else ''
            model = str(row[5]).strip() if len(row) > 5 and row[5] else ''
            serial_number = str(row[6]).strip() if len(row) > 6 and row[6] else ''
            assignee = str(row[8]).strip() if len(row) > 8 and row[8] else ''
            department = str(row[13]).strip() if len(row) > 13 and row[13] else ''
            
            order_number = str(row[38]).strip() if len(row) > 38 and row[38] else ''
            disk = str(row[22]).strip() if len(row) > 22 and row[22] else ''
            memory = str(row[23]).strip() if len(row) > 23 and row[23] else ''
            cpu = str(row[24]).strip() if len(row) > 24 and row[24] else ''
            oa_payment_number = str(row[37]).strip() if len(row) > 37 and row[37] else ''

            price_val = row[35] if len(row) > 35 else 0
            try:
                price = float(str(price_val).replace(',', '').replace('¥', '').replace('元', '').strip()) if price_val else 0.0
            except:
                price = 0.0
                
            date_val = row[16] if len(row) > 16 else ''
            purchase_date = ''
            if isinstance(date_val, datetime.datetime):
                purchase_date = date_val.strftime('%Y-%m-%d')
            elif date_val:
                purchase_date = str(date_val).strip()

            cursor = conn.execute("SELECT id FROM assets WHERE asset_id = ?", (asset_id,))
            if cursor.fetchone():
                continue
                
            conn.execute('''
                INSERT INTO assets (asset_id, name, asset_type, brand, model, serial_number, assignee, department, purchase_date, purchase_price, status, created_at, updated_at, order_number, disk, memory, cpu, oa_payment_number)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ''', (asset_id, name, asset_type, brand, model, serial_number, assignee, department, purchase_date, price, status, timestamp, timestamp, order_number, disk, memory, cpu, oa_payment_number))
            
            details = json.dumps({"action": "Excel批量导入", "file": filename}, ensure_ascii=False)
            conn.execute('''
                INSERT INTO operation_log (timestamp, asset_id, operation, operator, details)
                VALUES (?, ?, 'IMPORT', 'System', ?)
            ''', (timestamp, asset_id, details))
            
            imported_count += 1
            
    conn.commit()
    conn.close()
    return imported_count

@app.post("/api/upload_excel")
async def upload_excel(request: Request, file: UploadFile = File(...)):
    if not file.filename or not file.filename.endswith('.xlsx'):
        return RedirectResponse(url="/?error=not_excel", status_code=303)
        
    tmp_path = None
    try:
        with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp:
            content = await file.read()
            tmp.write(content)
            tmp_path = tmp.name

        imported_count = parse_excel_and_insert(tmp_path, file.filename)
        log.info(f"Excel导入成功: {file.filename}, 导入{imported_count}条")
        return RedirectResponse(url=f"/?success=imported_{imported_count}", status_code=303)
    except Exception as e:
        log.error(f"Excel导入失败: {e}", exc_info=True)
        return RedirectResponse(url=f"/?error=import_failed&detail={str(e)[:100]}", status_code=303)
    finally:
        if tmp_path and os.path.exists(tmp_path):
            try:
                os.unlink(tmp_path)
            except:
                pass

@app.get("/api/export_excel")
async def export_excel():
    """导出Excel，懒导入openpyxl"""
    try:
        import openpyxl
    except ImportError:
        return HTMLResponse(content="<p style='color:red;'>缺少 openpyxl 依赖，请运行：pip install openpyxl</p>", status_code=500)
    
    conn = None
    try:
        conn = get_db()
        cursor = conn.execute("SELECT * FROM assets ORDER BY REPLACE(purchase_date,'/','-') DESC, asset_id DESC")
        assets = cursor.fetchall()
    finally:
        if conn:
            conn.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "资产台账"
    
    headers = ["序号", "资产编码", "资产名称", "资产分类", "品牌", "型号", "设备序列号", "使用人", "使用部门", "采购日期", "采购价格(¥)", "状态"]
    ws.append(headers)
    
    status_map = {"in_use": "在用", "in_stock": "空闲", "scrapped": "已报废"}
    type_map_zh_xls = {'laptop': '笔记本', 'desktop': '台式机', 'monitor': '显示器', 'phone': '手机', 'tablet': '平板', 'keyboard_mouse': '键鼠', 'printer': '打印机', 'other': '其他'}
    
    for idx, a in enumerate(assets, 1):
        ws.append([
            idx, a["asset_id"], a["name"], type_map_zh_xls.get(a["asset_type"], a["asset_type"]),
            a["brand"], a["model"], a["serial_number"], a["assignee"], a["department"],
            a["purchase_date"], a["purchase_price"], status_map.get(a["status"], a["status"])
        ])
        
    with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp:
        wb.save(tmp.name)
        tmp_path = tmp.name
        
    return FileResponse(
        path=tmp_path, 
        filename=f"IT 资产管理系统资产台账_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
        media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

# ── 日志格式化 ──────────────────────────────────────────────────────────────────
def format_log_details(operation, details_json):
    """将操作日志的details JSON格式化为可读的中文描述"""
    try:
        d = json.loads(details_json) if isinstance(details_json, str) else (details_json or {})
    except (json.JSONDecodeError, TypeError):
        d = {}
    
    if operation == "退库入库":
        return "资产退库入库，使用人已清空，转入闲置库存"
    elif operation == "STATUS_CHANGE":
        old = d.get("from", d.get("old_status", "?"))
        new = d.get("to", d.get("new_status", "?"))
        status_map = {"in_use": "分配使用", "in_stock": "闲置入库", "scrapped": "报废", "disposed": "处置"}
        return f"状态变更: {status_map.get(old, old)} → {status_map.get(new, new)}"
    elif operation == "INVENTORY":
        loc = d.get("location", "现场")
        return f"扫码盘点确认（地点: {loc}）"
    elif operation == "IMPORT":
        fn = d.get("file", "")
        return f"系统数据导入（{fn}）" if fn else "系统数据导入"
    elif operation == "ASSIGN":
        return f"分配至: {d.get('assignee', '?')}"
    else:
        return d.get("action", operation)

@app.get("/api/scan/{asset_id}", response_class=HTMLResponse)
async def scan_asset(request: Request, asset_id: str):
    conn = None
    try:
        conn = get_db()
        cursor = conn.execute("SELECT * FROM assets WHERE asset_id = ?", (asset_id,))
        asset = cursor.fetchone()
        
        cursor_log = conn.execute("SELECT * FROM operation_log WHERE asset_id = ? ORDER BY timestamp DESC", (asset_id,))
        logs = [dict(row) for row in cursor_log.fetchall()]
        
        if not asset:
            return HTMLResponse(
                content=f"""<html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"></head>
                <body style="text-align:center;padding:50px;font-family:sans-serif;">
                    <h2>未找到资产</h2><p style="color:#666;">编码: {asset_id}</p>
                    <a href="/" style="color:#3498db;">返回首页</a>
                </body></html>""",
                status_code=404
            )
            
        # 为每条日志生成可读的中文摘要
        for log in logs:
            log["summary"] = format_log_details(log["operation"], log.get("details", ""))
        
        return templates.TemplateResponse(request, "scan.html", {"request": request, "asset": asset, "logs": logs})
    except Exception as e:
        log.error(f"扫码页加载失败 [{asset_id}]: {e}")
        return HTMLResponse(content=f"<p style='color:red;'>加载失败: {e}</p>", status_code=500)
    finally:
        if conn:
            try:
                conn.close()
            except:
                pass

@app.get("/asset/{asset_id}", response_class=HTMLResponse)
async def asset_detail_page(request: Request, asset_id: str):
    """资产档案详情页（PC管理视角，便于审计）：全字段 + 审计履历 + 可改状态/使用人"""
    conn = None
    try:
        conn = get_db()
        cursor = conn.execute("SELECT * FROM assets WHERE asset_id = ?", (asset_id,))
        asset = cursor.fetchone()
        if not asset:
            return HTMLResponse(
                content=f"""<html><head><meta charset="utf-8"></head><body style="text-align:center;padding:80px;font-family:sans-serif;"><h2>未找到资产</h2><p style="color:#666;">编码: {asset_id}</p><a href="/" style="color:#3498db;">返回首页</a></body></html>""",
                status_code=404)
        cursor_log = conn.execute("SELECT * FROM operation_log WHERE asset_id = ? ORDER BY timestamp DESC", (asset_id,))
        logs = [dict(row) for row in cursor_log.fetchall()]
        for log_ in logs:
            log_["summary"] = format_log_details(log_.get("operation", ""), log_.get("details", ""))
        return templates.TemplateResponse(request, "asset_detail_page.html", {"request": request, "asset": asset, "logs": logs})
    except Exception as e:
        log.error(f"资产详情页加载失败 [{asset_id}]: {e}")
        return HTMLResponse(content=f"<p style='color:red;'>加载失败: {e}</p>", status_code=500)
    finally:
        if conn:
            try:
                conn.close()
            except:
                pass


@app.get("/api/print_label/{asset_id}", response_class=HTMLResponse)
async def print_label(request: Request, asset_id: str):
    """补打标签：生成单张精臣B50标签页，支持浏览器直接打印"""
    conn = None
    try:
        conn = get_db()
        cursor = conn.execute("SELECT * FROM assets WHERE asset_id = ?", (asset_id,))
        asset = cursor.fetchone()
        if not asset:
            return HTMLResponse(
                content=f"""<html><head><meta charset="utf-8"></head>
                <body style="text-align:center;padding:50px;font-family:sans-serif;">
                    <h3>未找到资产: {asset_id}</h3>
                    <a href="/" style="color:#3498db;">返回首页</a>
                </body></html>""",
                status_code=404
            )
        
        # ── 生成二维码 ──────────────────────────────────────────────────────────
        # 根治架构：二维码指向固定公网入口（公司域名），后端 web_server /api/scan 实时查库，标签永不失效
        # 域名通过环境变量 ASSET_PUBLIC_URL 注入（IT映射后填真实域名），缺省回退占位域名
        # 兼容 CloudStudio 临时公网方案：若地址含 codebuddy.work，则用 ?aid= 格式（静态页形态）
        import os as _os
        # 固定公网入口：默认指向 GitHub Pages 永久地址（标签二维码根治方案，地址永不变）
        PUBLIC_EXTERNAL_URL = _os.environ.get(
            "ASSET_PUBLIC_URL",
            "https://YOUR_USERNAME.github.io/YOUR_REPO/asset-labels"
        ).rstrip("/")
        # 静态页形态（CloudStudio / GitHub Pages）：二维码用 ?aid= 格式，地址永久不变
        if "codebuddy.work" in PUBLIC_EXTERNAL_URL or "github.io" in PUBLIC_EXTERNAL_URL:
            scan_url = f"{PUBLIC_EXTERNAL_URL}/asset_detail.html?aid={asset_id}"
        else:
            # 后端实时查库形态（公司固定域名 + web_server /api/scan）
            scan_url = f"{PUBLIC_EXTERNAL_URL}/api/scan/{asset_id}"
        # 生成PNG二维码 → base64 → 直接内嵌为<img src="data:image/png;base64,...">
        qr = qrcode.QRCode(box_size=8, border=1)
        qr.add_data(scan_url)
        qr.make(fit=True)
        buf = io.BytesIO()
        qr.make_image(fill_color="black", back_color="white").save(buf, format="PNG")
        qr_b64 = base64.b64encode(buf.getvalue()).decode("utf-8")
        qr_data_uri = f"data:image/png;base64,{qr_b64}"
        
        return templates.TemplateResponse(request, "print_label.html", {"request": request, "asset": asset, "qr_data_uri": qr_data_uri})
    except Exception as e:
        log.error(f"补打标签失败 [{asset_id}]: {e}")
        return HTMLResponse(content=f"<p style='color:red;'>加载失败: {e}</p>", status_code=500)
    finally:
        if conn:
            try: conn.close()
            except: pass

@app.post("/api/scan_confirm")
async def scan_confirm(request: Request, asset_id: str = Form(...), operator: str = Form(...)):
    conn = None
    try:
        conn = get_db()
        timestamp = datetime.datetime.now().isoformat()
        details = json.dumps({"action": "扫码盘点", "location": "现场"}, ensure_ascii=False)
        conn.execute('''
            INSERT INTO operation_log (timestamp, asset_id, operation, operator, details)
            VALUES (?, ?, 'INVENTORY', ?, ?)
        ''', (timestamp, asset_id, operator, details))
        conn.commit()
        log.info(f"扫码盘点确认: {asset_id} by {operator}")
    except Exception as e:
        log.error(f"扫码盘点失败 [{asset_id}]: {e}")
    finally:
        if conn:
            try:
                conn.close()
            except:
                pass
    return RedirectResponse(url=f"/api/scan/{asset_id}?success=1", status_code=303)

if __name__ == "__main__":
    import uvicorn
    port = int(os.environ.get("PORT", "8081"))
    log.info("=" * 60)
    log.info(f"资产大盘服务启动: http://0.0.0.0:{port}")
    log.info(f"数据库路径: {DB_PATH}")
    log.info(f"日志路径: {LOG_FILE}")
    log.info("=" * 60)
    uvicorn.run(app, host="0.0.0.0", port=port, log_config=None)  # 使用自定义日志配置
