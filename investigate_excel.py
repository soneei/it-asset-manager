import openpyxl

EXCEL_PATH = 'input.xlsx'

def check_all_sheets():
    wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    total_assets = 0
    print(f"All sheets: {wb.sheetnames}")
    for sheet_name in wb.sheetnames:
        if sheet_name == '汇总_整表大盘':
            continue
        ws = wb[sheet_name]
        count = 0
        missing_id_count = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            if any(row):  # If the row is not completely empty
                if row[1]: # has asset_id
                    count += 1
                else:
                    missing_id_count += 1
        print(f"Sheet: {sheet_name}, valid IDs: {count}, missing IDs: {missing_id_count}")
        total_assets += count + missing_id_count
    print(f"Total potential assets across all detailed sheets: {total_assets}")

if __name__ == "__main__":
    check_all_sheets()