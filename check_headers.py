import openpyxl

EXCEL_PATH = 'input.xlsx'
wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)
for sheet in wb.sheetnames:
    if '在用' in sheet or '空闲' in sheet:
        ws = wb[sheet]
        rows = list(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        if rows:
            print(f"Sheet: {sheet}")
            for i, col in enumerate(rows[0]):
                if col:
                    print(f"{i}: {col}")
        break
