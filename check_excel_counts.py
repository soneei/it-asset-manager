import openpyxl

EXCEL_PATH = 'input.xlsx'

def check_excel_counts():
    wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    total_valid_rows = 0
    unique_ids = set()
    sheet_counts = {}
    
    for sheet_name in wb.sheetnames:
        if '在用' in sheet_name or '空闲' in sheet_name:
            ws = wb[sheet_name]
            count = 0
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row and len(row) > 1 and row[1]:
                    count += 1
                    unique_ids.add(str(row[1]).strip())
            sheet_counts[sheet_name] = count
            total_valid_rows += count
            print(f"Sheet '{sheet_name}' 包含 {count} 条数据")
            
    print(f"\n所有 sheet 累计有效行数: {total_valid_rows}")
    print(f"全局唯一的资产编码 (Asset ID) 数量: {len(unique_ids)}")

if __name__ == "__main__":
    check_excel_counts()