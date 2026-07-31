#!/usr/bin/env python3
"""
从Excel资产清单读取两个sheet数据，生成JSON并重新生成仪表盘
"""
import json, os, datetime
import openpyxl

EXCEL_PATH = 'input.xlsx'
OUTPUT_JSON = os.path.join(os.path.dirname(__file__), 'data', 'assets.json')
OUTPUT_HTML = os.path.join(os.path.dirname(__file__), '..', 'dashboard.html')

TODAY = datetime.date(2026, 5, 19)

# 类型映射：从Excel中文分类映射到英文
TYPE_MAP = {
    '台式机': 'desktop', '笔记本': 'laptop', '显示器': 'monitor',
    '手机': 'phone', '平板': 'tablet', '键鼠': 'keyboard_mouse',
    '打印机': 'printer', '其他': 'other',
    '一体机': 'desktop', '电脑主机': 'desktop',
}

def parse_price(val):
    if val is None:
        return 0
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).replace(',', '').replace('¥', '').replace('元', '').strip()
    try:
        return float(s) if s else 0
    except:
        return 0

def parse_date(val):
    if val is None:
        return ''
    if isinstance(val, datetime.datetime):
        return val.strftime('%Y-%m-%d')
    s = str(val).strip()
    if not s or s == 'None':
        return ''
    return s

def calc_depreciation(price, purchase_date, dep_years=3):
    """直线折旧，残值率5%，返回(原值, 当前净值)"""
    if not price or not purchase_date:
        return price or 0, price or 0
    try:
        if '/' in purchase_date:
            dt = datetime.datetime.strptime(purchase_date, '%Y/%m/%d').date()
        elif '-' in purchase_date:
            dt = datetime.datetime.strptime(purchase_date[:10], '%Y-%m-%d').date()
        else:
            return price, price
    except:
        return price, price
    
    years = (TODAY - dt).days / 365.25
    salvage = price * 0.05
    depreciable = price - salvage
    current = max(salvage, price - depreciable * min(years / dep_years, 1.0))
    return round(price, 2), round(current, 2)

def read_sheet(ws, status):
    """读取一个sheet的数据"""
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assets = []
    for row in rows:
        if not row or not row[1]:  # 编码为空则跳过
            continue
        
        asset_type_raw = str(row[3] or '').strip()
        asset_type = 'other'
        for k, v in TYPE_MAP.items():
            if k in asset_type_raw:
                asset_type = v
                break
        
        price = parse_price(row[35])  # 采购价
        date = parse_date(row[16])   # 采购日期
        
        # 折旧年限：默认3年
        dep_years = 3
        
        asset = {
            'asset_id': str(row[1]).strip(),
            'name': str(row[2] or '').strip(),
            'asset_type': asset_type,
            'brand': str(row[4] or '').strip(),
            'model': str(row[5] or '').strip(),
            'serial_number': str(row[6] or '').strip(),
            'status': status,
            'assignee': str(row[7] or '').strip(),  # 使用人/位置
            'department': str(row[13] or '').strip() if len(row) > 13 else '',  # 部门
            'location': str(row[7] or '').strip(),
            'purchase_date': date,
            'purchase_price': price,
            'depreciation_years': dep_years,
            'supplier': str(row[39] or '').strip() if len(row) > 39 else '',
            'warranty_expire': str(row[22] or '').strip() if len(row) > 22 else '',
            'remarks': str(row[30] or '').strip() if len(row) > 30 else '',
            'disposal_amount': parse_price(row[37]) if len(row) > 37 else 0,
        }
        assets.append(asset)
    return assets

# 读取Excel
wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)
ws1 = wb[wb.sheetnames[0]]  # 在用
ws2 = wb[wb.sheetnames[1]]  # 空闲

in_use_list = read_sheet(ws1, 'in_use')
idle_list = read_sheet(ws2, 'idle')
all_assets = in_use_list + idle_list

print(f'Excel读取: 在用 {len(in_use_list)} 条, 空闲 {len(idle_list)} 条, 总计 {len(all_assets)} 条')

# 保存JSON
data = {
    'version': '1.0',
    'exported_at': datetime.datetime.now().isoformat(),
    'total_count': len(all_assets),
    'assets': all_assets,
}
with open(OUTPUT_JSON, 'w', encoding='utf-8') as f:
    json.dump(data, f, ensure_ascii=False, indent=2)
print(f'JSON已保存: {OUTPUT_JSON}')

# ========== 生成仪表盘 ==========

TYPE_LABELS = {
    'laptop': '笔记本', 'desktop': '台式机', 'monitor': '显示器',
    'phone': '手机', 'tablet': '平板', 'keyboard_mouse': '键鼠',
    'printer': '打印机', 'other': '其他'
}

def fmt_price(v):
    if not v or v == 0:
        return '-'
    return f'¥{int(v):,}'

def safe(v):
    if v is None or v == '' or v == 0.0:
        return '-'
    return str(v)

def price_cell(orig, curr):
    if not orig:
        return '-'
    if curr and curr != orig:
        return f'<span style="color:#ccc;text-decoration:line-through;font-size:10px;">{fmt_price(orig)}</span> <span style="color:#3b82f6;font-weight:600;">{fmt_price(curr)}</span>'
    return fmt_price(orig)

# 统计
iu_orig = sum(a['purchase_price'] for a in in_use_list)
idle_orig = sum(a['purchase_price'] for a in idle_list)

# 计算折旧
for a in all_assets:
    a['original_price'], a['current_value'] = calc_depreciation(
        a['purchase_price'], a['purchase_date'], a['depreciation_years']
    )

iu_current = sum(a['current_value'] for a in in_use_list)
idle_current = sum(a['current_value'] for a in idle_list)
total_orig = iu_orig + idle_orig
total_current = iu_current + idle_current

# 处置统计
disposed_list = [a for a in all_assets if a['status'] == 'disposed']
disp_total = sum(a['disposal_amount'] for a in disposed_list)
disp_book = sum(a['current_value'] for a in disposed_list)

# 员工分组
from collections import defaultdict
emp_assets = defaultdict(list)
for a in in_use_list:
    emp = a.get('assignee') or '(未分配)'
    emp_assets[emp].append(a)

# CSS
CSS = """
*{margin:0;padding:0;box-sizing:border-box;}body{font-family:-apple-system,BlinkMacSystemFont,'Segoe UI',Roboto,'PingFang SC','Microsoft YaHei',sans-serif;background:#f0f2f5;color:#333;font-size:12px;line-height:1.5;}.container{max-width:100%;padding:16px;}.top-bar{background:linear-gradient(135deg,#1e3a5f,#2d5a87);color:#fff;padding:18px 24px;border-radius:12px;display:flex;justify-content:space-between;align-items:center;margin-bottom:16px;flex-wrap:wrap;gap:12px;}.top-bar h1{font-size:20px;font-weight:700;}.top-bar .sub{opacity:0.75;font-size:11px;margin-top:2px;}.stats{display:flex;gap:20px;flex-wrap:wrap;}.stat{text-align:center;min-width:90px;}.stat .num{font-size:20px;font-weight:700;}.stat .txt{font-size:10px;opacity:0.7;}.tag{display:inline-block;padding:2px 8px;border-radius:10px;font-size:10px;font-weight:600;margin-left:3px;}.tag-blue{background:rgba(59,130,246,0.3);}.tag-orange{background:rgba(245,158,11,0.3);}
.tab-bar{display:flex;gap:4px;margin-bottom:12px;}.tab-btn{padding:10px 28px;border:none;border-radius:10px 10px 0 0;font-size:14px;font-weight:600;cursor:pointer;background:#dcdfe6;color:#909399;transition:all .25s;}.tab-btn.active{background:#fff;color:#1e3a5f;box-shadow:0 -3px 12px rgba(30,58,95,.08);}.panel{background:#fff;border-radius:0 12px 12px 12px;box-shadow:0 2px 12px rgba(0,0,0,.06);overflow:hidden;display:none;}.panel.active{display:block;}
.filter-bar{padding:10px 16px;background:#f8f9fb;border-bottom:1px solid #eee;display:flex;gap:8px;align-items:center;flex-wrap:wrap;}.filter-label{font-weight:600;color:#555;font-size:11px;margin-right:4px;}.f-btn{padding:4px 12px;border:1px solid #ddd;border-radius:6px;background:#fff;font-size:11px;cursor:pointer;}.f-btn.active{background:#1e3a5f;color:#fff;border-color:#1e3a5f;}.f-btn:hover:not(.active){border-color:#1e3a5f;color:#1e3a5f;}.search-box{margin-left:auto;display:flex;align-items:center;gap:4px;}.search-input{padding:4px 10px;border:1px solid #ddd;border-radius:6px;font-size:11px;width:180px;outline:none;}.search-input:focus{border-color:#3b82f6;}
.table-wrap{overflow-x:auto;}table{width:100%;border-collapse:collapse;min-width:1500px;}th{background:linear-gradient(#f7f9fc,#edf1f6);padding:8px 7px;text-align:left;font-weight:600;color:#555;font-size:11px;border-bottom:2px solid #dde2eb;position:sticky;top:0;z-index:2;}td{padding:6px 7px;border-bottom:1px solid #f0f0f0;vertical-align:middle;font-size:11px;color:#444;}tr:hover td{background:#f5f8ff;}
.emp-header{background:#f0f4ff!important;cursor:pointer;}.emp-header td{font-weight:600;color:#1e40af;border-bottom:1px solid #c7d6fa;padding:8px 10px;}.emp-header:hover td{background:#e6ecff;}.emp-arrow{display:inline-block;width:0;height:0;border-left:5px solid transparent;border-right:5px solid transparent;border-top:6px solid #3b82f6;margin-right:6px;transition:transform .2s;vertical-align:middle;}.collapsed .emp-arrow{transform:rotate(-90deg);}.chips{display:inline-flex;gap:4px;margin-left:8px;}.chip{padding:1px 7px;border-radius:10px;font-size:10px;font-weight:500;}.chip-laptop{background:#dbeafe;color:#1e40af;}.chip-desktop{background:#d1fae5;color:#065f46;}.chip-monitor{background:#fef3c7;color:#92400e;}
.detail-row{display:none;}.detail-row.show{display:table-row;}.detail-row td{padding-left:32px;background:#fafbff;}
.st{display:inline-flex;align-items:center;padding:2px 9px;border-radius:10px;font-size:10px;font-weight:500;}.st-idle{background:#fef3c7;color:#92400e;}.st-disposed{background:#fee2e2;color:#991b1b;}.st-in_stock{background:#dbeafe;color:#1e40af;}
.code{font-family:'Monaco','Consolas',monospace;font-size:10px;background:#f5f5f5;padding:2px 6px;border-radius:4px;color:#666;}
.idle-tr td{background:#fffef8;}.idle-tr:nth-child(even) td{background:#fffdf5;}.idle-tr:hover td{background:#fff8e1;}
.dcard{background:#fff;padding:10px 16px;border-radius:8px;min-width:140px;box-shadow:0 1px 4px rgba(0,0,0,.06);}.dlbl{font-size:10px;color:#999;margin-bottom:4px;}.dval{font-size:16px;font-weight:700;}.dval.blue{color:#3b82f6;}.dval.orange{color:#f59e0b;}.dval.red{color:#ef4444;}
.footer{text-align:center;color:#bbb;font-size:11px;padding:16px;}
"""

JS = """
function switchTab(id,btn){document.querySelectorAll('.panel').forEach(p=>p.classList.remove('active'));document.querySelectorAll('.tab-btn').forEach(b=>b.classList.remove('active'));document.getElementById('panel-'+id).classList.add('active');btn.classList.add('active');}
function toggleEmp(eid,row){row.classList.toggle('collapsed');document.querySelectorAll('[data-parent="'+eid+'"]').forEach(r=>r.classList.toggle('show'));}
function filterType(btn,type){document.querySelectorAll('#panel-in-use .f-btn').forEach(b=>b.classList.remove('active'));btn.classList.add('active');document.querySelectorAll('#in-use-body .emp-header').forEach(r=>{r.style.display=(type==='all'||r.dataset.types.includes(type))?'':'none';});document.querySelectorAll('#in-use-body .detail-row').forEach(r=>{r.style.display=(type==='all'||r.dataset.type===type)?'':'none';});}
function searchInUse(q){q=q.toLowerCase().trim();document.querySelectorAll('#in-use-body .emp-header').forEach(r=>{r.style.display=(!q||r.dataset.emp.toLowerCase().includes(q))?'':'none';});document.querySelectorAll('#in-use-body .detail-row').forEach(r=>{r.style.display=(!q||r.dataset.search.includes(q))?'':'none';});}
function filterIdle(btn,status){document.querySelectorAll('#panel-idle .f-btn').forEach(b=>b.classList.remove('active'));btn.classList.add('active');document.querySelectorAll('.idle-tr').forEach(tr=>{tr.style.display=(status==='all'||tr.dataset.status===status)?'':'none';});}
function searchIdle(q){q=q.toLowerCase().trim();document.querySelectorAll('.idle-tr').forEach(tr=>{tr.style.display=(!q||tr.dataset.search.includes(q))?'':'none';});}
"""

parts = []
A = parts.append

A(f'<html lang="zh-CN"><head><meta charset="UTF-8"><title>IT 资产管理系统资产台账</title>')
A(f'<style>{CSS}</style></head><body><div class="container">')

# 顶部栏
A(f'<div class="top-bar"><div><h1>IT 资产管理系统 IT 资产台账</h1>')
A(f'<div class="sub">数据截止 2026-05-19 · 折旧基准日 2026-05-19 · 直线折旧残值率5%</div></div>')
A(f'<div class="stats">')
A(f'<div class="stat"><div class="num">{len(all_assets)}</div><div class="txt">资产总数</div></div>')
A(f'<div class="stat"><div class="num">{len(in_use_list)}<span class="tag tag-blue">在用</span></div><div class="txt">原值{fmt_price(iu_orig)}<br>现值{fmt_price(iu_current)}</div></div>')
A(f'<div class="stat"><div class="num">{len(idle_list)}<span class="tag tag-orange">空闲</span></div><div class="txt">原值{fmt_price(idle_orig)}<br>现值{fmt_price(idle_current)}</div></div>')
A(f'<div class="stat"><div class="num">{fmt_price(total_current)}</div><div class="txt">总净值（原值{fmt_price(total_orig)}）</div></div>')
if disposed_list:
    A(f'<div class="stat"><div class="num">{len(disposed_list)}<span class="tag tag-red">已处置</span></div><div class="txt">处置金额{fmt_price(disp_total)}</div></div>')
A('</div></div>')

# Tabs
tabs = [('in-use', f'① 在用资产（{len(in_use_list)}）'), ('idle', f'② 空闲资产（{len(idle_list)}）')]
if disposed_list:
    tabs.append(('disposed', f'③ 处置资产（{len(disposed_list)}）'))
A('<div class="tab-bar">')
for i, (tid, lbl) in enumerate(tabs):
    act = ' active' if i == 0 else ''
    A(f'<button class="tab-btn{act}" onclick="switchTab(\'{tid}\',this)">{lbl}</button>')
A('</div>')

# Panel 1: 在用
A('<div class="panel active" id="panel-in-use">')
A('<div class="filter-bar"><span class="filter-label">筛选：</span>')
A('<button class="f-btn active" onclick="filterType(this,\'all\')">全部</button>')
A('<button class="f-btn" onclick="filterType(this,\'laptop\')">笔记本</button>')
A('<button class="f-btn" onclick="filterType(this,\'desktop\')">台式机</button>')
A('<button class="f-btn" onclick="filterType(this,\'monitor\')">显示器</button>')
A('<div class="search-box"><input type="text" class="search-input" placeholder="搜索员工/编码/名称..." onkeyup="searchInUse(this.value)"></div></div>')
A('<div class="table-wrap"><table>')
A('<thead><tr><th style="width:200px">员工</th><th style="width:140px">设备汇总</th><th style="width:115px">编码</th><th>设备名称</th><th>分类</th><th>品牌</th><th>型号</th><th>采购价</th><th>当前净值</th><th>采购日期</th><th>部门</th><th>位置</th><th>备注</th></tr></thead>')
A('<tbody id="in-use-body">')

for emp_name, alist in sorted(emp_assets.items(), key=lambda x: -len(x[1])):
    if emp_name == '(未分配)':
        continue
    chips = []
    for t in ['laptop','desktop','monitor','phone','tablet']:
        cnt = sum(1 for a in alist if a['asset_type']==t)
        if cnt:
            c = {'laptop':'chip-laptop','desktop':'chip-desktop','monitor':'chip-monitor'}.get(t,'')
            chips.append(f'<span class="chip {c}">{TYPE_LABELS.get(t,t)}×{cnt}</span>')
    eid = emp_name.replace(' ','_').replace('(','').replace(')','')
    types_all = ' '.join(a['asset_type'] for a in alist)
    A(f'<tr class="emp-header" onclick="toggleEmp(\'{eid}\',this)" data-emp="{emp_name}" data-types="{types_all}">')
    A(f'<td><span class="emp-arrow"></span>{emp_name}</td>')
    A(f'<td><span class="chips">{"".join(chips)}</span></td><td colspan="11"></td></tr>')
    
    for i, a in enumerate(alist):
        t_label = TYPE_LABELS.get(a['asset_type'], '-')
        show = ' show' if i == 0 else ''
        pc = price_cell(a['original_price'], a['current_value'])
        sstr = (emp_name + ' ' + a['asset_id'] + ' ' + a['name']).lower()
        A(f'<tr class="detail-row{show}" data-parent="{eid}" data-type="{a["asset_type"]}" data-search="{sstr}">')
        A(f'<td></td><td></td>')
        A(f'<td><span class="code">{a["asset_id"]}</span></td>')
        A(f'<td>{safe(a["name"])}</td><td>{t_label}</td><td>{safe(a["brand"])}</td><td>{safe(a["model"])}</td>')
        A(f'<td>{pc}</td>')
        A(f'<td>{safe(a["purchase_date"])}</td><td>{safe(a["department"])}</td><td>{safe(a["location"])}</td><td style="color:#999;font-size:10px;">{safe(a["remarks"])}</td>')
        A('</tr>')

unasgn = emp_assets.get('(未分配)', [])
if unasgn:
    A('<tr><td colspan="13" style="background:#fff8e6;font-weight:600;color:#92400e;padding:10px;">⚠️ 未分配人员</td></tr>')
    for a in unasgn:
        t_label = TYPE_LABELS.get(a['asset_type'], '-')
        A(f'<tr><td></td><td></td><td><span class="code">{a["asset_id"]}</span></td><td>{safe(a["name"])}</td><td>{t_label}</td><td>{safe(a["brand"])}</td><td>{safe(a["model"])}</td><td>{price_cell(a["original_price"],a["current_value"])}</td><td>{safe(a["purchase_date"])}</td><td>{safe(a["department"])}</td><td>{safe(a["location"])}</td><td>{safe(a["remarks"])}</td></tr>')

A('</tbody></table></div></div>')

# Panel 2: 空闲
A(f'<div class="panel" id="panel-idle">')
A(f'<div class="filter-bar"><span class="filter-label">状态：</span>')
A(f'<button class="f-btn active" onclick="filterIdle(this,\'all\')">全部（{len(idle_list)}）</button>')
for st in ['idle','disposed','in_stock','under_repair']:
    cnt = sum(1 for a in idle_list if a['status']==st)
    if cnt:
        A(f'<button class="f-btn" onclick="filterIdle(this,\'{st}\')">{st}（{cnt}）</button>')
A(f'<div class="search-box"><input type="text" class="search-input" placeholder="搜索编码/名称..." onkeyup="searchIdle(this.value)"></div></div>')
A(f'<div class="table-wrap"><table>')
A(f'<thead><tr><th style="width:50px">#</th><th style="width:90px">状态</th><th style="width:115px">编码</th><th>设备名称</th><th>分类</th><th>品牌</th><th>型号</th><th>采购价</th><th>当前净值</th><th>采购日期</th><th>部门</th><th>位置</th><th>供应商</th><th>备注</th></tr></thead><tbody>')

for idx, a in enumerate(idle_list, 1):
    st_cls = {'idle':'st-idle','disposed':'st-disposed','in_stock':'st-in_stock'}.get(a['status'],'')
    t_label = TYPE_LABELS.get(a['asset_type'], '-')
    pc = price_cell(a['original_price'], a['current_value'])
    sstr = (a['asset_id'] + ' ' + a['name'] + ' ' + a.get('supplier','')).lower()
    A(f'<tr class="idle-tr" data-status="{a["status"]}" data-search="{sstr}">')
    A(f'<td style="color:#bbb;text-align:center;">{idx}</td>')
    A(f'<td><span class="st {st_cls}">{a["status"]}</span></td>')
    A(f'<td><span class="code">{a["asset_id"]}</span></td>')
    A(f'<td>{safe(a["name"])}</td><td>{t_label}</td><td>{safe(a["brand"])}</td><td>{safe(a["model"])}</td>')
    A(f'<td>{pc}</td>')
    A(f'<td>{safe(a["purchase_date"])}</td><td>{safe(a["department"])}</td><td>{safe(a["location"])}</td><td>{safe(a["supplier"])}</td><td style="color:#999;font-size:10px;">{safe(a["remarks"])}</td>')
    A('</tr>')

A('</tbody></table></div></div>')

# Panel 3: 处置
if disposed_list:
    A('<div class="panel" id="panel-disposed">')
    A('<div class="filter-bar" style="background:#fff8f0;">')
    A(f'<div class="dcard"><div class="dlbl">处置数量</div><div class="dval">{len(disposed_list)} 台</div></div>')
    A(f'<div class="dcard"><div class="dlbl">账面净值</div><div class="dval orange">{fmt_price(disp_book)}</div></div>')
    A(f'<div class="dcard"><div class="dlbl">处置金额</div><div class="dval blue">{fmt_price(disp_total)}</div></div>')
    gain = disp_total - disp_book
    A(f'<div class="dcard"><div class="dlbl">处置损益</div><div class="dval {"red" if gain<0 else "blue"}">{fmt_price(abs(gain))} {"损失" if gain<0 else "收益"}</div></div>')
    A('</div><div class="table-wrap"><table>')
    A('<thead><tr><th style="width:50px">#</th><th style="width:115px">编码</th><th>名称</th><th>分类</th><th>品牌/型号</th><th>原值</th><th>处置时净值</th><th>处置金额</th><th>处置损益</th><th>处置日期</th><th>备注</th></tr></thead><tbody>')
    for idx, a in enumerate(disposed_list, 1):
        t_label = TYPE_LABELS.get(a['asset_type'], '-')
        gain = a.get('disposal_amount',0) - a['current_value']
        gain_str = f'{"-" if gain<0 else "+"}{fmt_price(abs(gain))}' if gain != 0 else '-'
        A(f'<tr><td style="color:#bbb;text-align:center;">{idx}</td><td><span class="code">{a["asset_id"]}</span></td><td>{safe(a["name"])}</td><td>{t_label}</td><td>{safe(a["brand"])} {safe(a["model"])}</td><td>{fmt_price(a["purchase_price"])}</td><td><span style="color:#3b82f6;">{fmt_price(a["current_value"])}</span></td><td><span style="color:#10b981;font-weight:600;">{fmt_price(a["disposal_amount"])}</span></td><td>{gain_str}</td><td>-</td><td style="color:#999;font-size:10px;">{safe(a["remarks"])}</td></tr>')
    A('</tbody></table></div></div>')

A(f'</div><script>{JS}</script>')
A(f'<div class="footer">IT 资产管理系统 IT 资产台账 · 数据源自 input.xlsx · WorkBuddy</div>')
A('</div></body></html>')

with open(OUTPUT_HTML, 'w', encoding='utf-8') as f:
    f.write('\n'.join(parts))

print(f'\n仪表盘已生成: {OUTPUT_HTML}')
print(f'在用: {len(in_use_list)} 台  原值{fmt_price(iu_orig)} 现值{fmt_price(iu_current)}')
print(f'空闲: {len(idle_list)} 台  原值{fmt_price(idle_orig)} 现值{fmt_price(idle_current)}')
print(f'总净值: {fmt_price(total_current)} / 原值: {fmt_price(total_orig)}')
