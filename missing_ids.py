import openpyxl

EXCEL_PATH = 'input.xlsx'

def extract_missing_ids():
    wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    
    print("=== 缺少资产编码的记录清单 ===")
    
    for sheet_name in wb.sheetnames:
        if sheet_name == '汇总_整表大盘':
            continue
            
        ws = wb[sheet_name]
        
        # 获取表头以对应字段
        header = []
        for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
            header = [str(c).strip() if c else '' for c in row]
            
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if any(row):  # 行不全为空
                if not row[1]:  # 第二列(资产编码)为空
                    print(f"\n找到一条漏登编码的记录 -> 位置：Sheet '{sheet_name}', 第 {row_idx} 行")
                    print("具体信息:")
                    # 打印前几个关键列
                    for col_idx in [2, 3, 4, 5, 6, 8]: # 名称, 分类, 品牌, 型号, 序列号, 使用人
                        if col_idx < len(row):
                            col_name = header[col_idx] if col_idx < len(header) else f"列{col_idx+1}"
                            val = row[col_idx] if row[col_idx] else '(空)'
                            print(f"  - {col_name}: {val}")

if __name__ == "__main__":
    extract_missing_ids()